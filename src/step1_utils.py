"""
src/utils.py  |  STEP 1 — Utilities (Foundation)
------------------------------------------------
The very first building block. Every other module depends on this.
Contains:
  - clean_text()          : strips HTML, lowercases, normalises whitespace
  - extract_resume_text() : reads PDF and DOCX files into plain text
  - segment_concatenated_skills() : fixes run-together skill words from scraped sites
"""

import re
import io
import pdfplumber
from docx import Document


def clean_text(text: str) -> str:
    """Normalise raw text for skill extraction and embedding."""
    if not isinstance(text, str) or not text.strip():
        return ""
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", text)  # camelCase split
    text = re.sub(r"&\w+;", " ", text)                 # HTML entities
    text = re.sub(r"<[^>]+>", " ", text)               # HTML tags
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s\+\#\.\-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def extract_resume_text(uploaded_file) -> tuple[str, str | None]:
    """
    Extract raw text from an uploaded PDF or DOCX file.

    Returns:
        (text, error_msg) — if extraction fails, text is "" and error_msg explains why.
    """
    filename = uploaded_file.name.lower()

    # ------------------------------------------------------------------ PDF
    if filename.endswith(".pdf"):
        try:
            text_parts = []
            with pdfplumber.open(uploaded_file) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            text = "\n".join(text_parts).strip()
            if not text:
                return "", (
                    "⚠️ Couldn't extract text from this PDF — it may be a **scanned image** "
                    "rather than selectable text.\n\n"
                    "**Try one of these fixes:**\n"
                    "- Use [ILovePDF OCR](https://www.ilovepdf.com/ocr-pdf) to convert it first\n"
                    "- Try [Adobe Acrobat OCR](https://www.adobe.com/acrobat/online/pdf-to-word.html)\n"
                    "- Export from your original document as a text-based PDF"
                )
            return text, None
        except Exception as e:
            return "", f"Failed to read PDF: {e}"

    # ------------------------------------------------------------------ DOCX
    if filename.endswith(".docx"):
        try:
            file_bytes = uploaded_file.read()
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also grab text from tables inside the DOCX
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)
            text = "\n".join(paragraphs).strip()
            if not text:
                return "", "⚠️ The DOCX file appears to be empty or uses an unsupported format."
            return text, None
        except Exception as e:
            return "", f"Failed to read DOCX: {e}"

    return "", "⚠️ Unsupported file type. Please upload a PDF or DOCX file."


def segment_concatenated_skills(text: str, skill_extractor) -> str:
    """
    Some job sites (e.g. Naukri 'Key Skills' sections) list skills back-to-back
    with no space between them, e.g. "pythondata analysisdata analyticssql".
    This inserts a space around every known skill phrase found in the text,
    even mid-word, so the PhraseMatcher can pick them up individually.
    """
    forms = sorted(
        set(skill_extractor.surface_to_canonical.keys()), key=len, reverse=True
    )
    for form in forms:
        if len(form) < 3:
            continue
        text = re.sub(re.escape(form), f" {form} ", text)
    return re.sub(r"\s+", " ", text).strip()
